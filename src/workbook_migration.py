from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import shutil

from openpyxl import load_workbook

from excel_reader import ACTIVITY_SHEET_NAME, LEGACY_ACTIVITY_SHEET_NAME, normalize_activity_key
from activity_output_registry import ACTIVITY_OUTPUT_DEFINITIONS, normalize_output_type

ACTIVITY_HEADERS = [
    "Activity_Key",
    "theme",
    "activity_name",
    "focus_area",
    "activity_type",
    "activity_description",
    "vprp_plan",
    "pdi_indicator",
    "activity_for",
    "targeted_populace",
    "activity_nature",
    "major_head",
    "minor_head",
    "is_directly_funded_by_panchayat",
    "estimated_completion_year",
    "estimated_completion_month",
    "estimated_completion_days",
    "start_year",
    "start_month",
    "expected_beneficiary_general",
    "expected_beneficiary_sc",
    "expected_beneficiary_st",
    "indicative_unit_cost",
    "total_indicative_cost",
    "expected_results",
    "flagship_scheme",
    "select_supported_department",
    "estimated_total_cost",
    "activity_output_type",
    "Final Action",
    "status",
    "processed_time",
    "error_message",
]

TRAINING_HEADERS = [
    "Activity_Key",
    "Training Category",
    "Organized By",
    "Subject of Training",
    "Village",
    "Amount",
    "Total Trainees",
    "Total Duration",
]

INLINE_TRAINING_SOURCE_HEADERS = {
    "training_category": "Training Category",
    "training_organized_by": "Organized By",
    "training_organised_by": "Organized By",
    "training_subject": "Subject of Training",
    "training_village": "Village",
    "training_amount": "Amount",
    "training_total_trainees": "Total Trainees",
    "training_total_duration_days": "Total Duration",
}


@dataclass
class MigrationReport:
    backup_path: Path
    activity_rows: int
    training_rows: int
    migrated_columns: list[str]
    created_sheets: list[str]


class WorkbookMigrationError(RuntimeError):
    pass


def _normalize_header(header: str) -> str:
    return "_".join(str(header or "").strip().lower().replace("/", " ").replace("-", " ").split())


def migrate_input_workbook(workbook_path: str) -> MigrationReport:
    path = Path(workbook_path)
    if not path.exists():
        raise WorkbookMigrationError(f"Workbook not found: {path}")

    backup = path.with_name("input_backup_before_output_sheet_migration.xlsx")
    shutil.copy2(path, backup)

    wb = load_workbook(path)
    created_sheets: list[str] = []

    if ACTIVITY_SHEET_NAME in wb.sheetnames:
        source_sheet = wb[ACTIVITY_SHEET_NAME]
    elif LEGACY_ACTIVITY_SHEET_NAME in wb.sheetnames:
        source_sheet = wb[LEGACY_ACTIVITY_SHEET_NAME]
    else:
        raise WorkbookMigrationError(
            f"Activities source sheet not found. Expected '{ACTIVITY_SHEET_NAME}' or '{LEGACY_ACTIVITY_SHEET_NAME}'"
        )

    rows = list(source_sheet.iter_rows(values_only=True))
    if not rows:
        raise WorkbookMigrationError("Source activity sheet is empty")

    raw_headers = [str(h or "").strip() for h in rows[0]]
    normalized_headers = {_normalize_header(h): idx for idx, h in enumerate(raw_headers)}

    activity_key_col = normalized_headers.get("activity_key")
    if activity_key_col is None:
        activity_key_col = -1

    training_source_map: dict[str, int] = {}
    for src, _ in INLINE_TRAINING_SOURCE_HEADERS.items():
        if src in normalized_headers:
            training_source_map[src] = normalized_headers[src]

    activity_records: list[dict[str, str]] = []
    training_records: list[dict[str, str]] = []

    seen_activity_keys: set[str] = set()
    seen_training_keys: set[str] = set()

    for row_idx, row in enumerate(rows[1:], start=1):
        row_values = ["" if v is None else str(v).strip() for v in row]

        key = ""
        if activity_key_col >= 0 and activity_key_col < len(row_values):
            key = normalize_activity_key(row_values[activity_key_col])
        if not key:
            key = f"ACT-{row_idx:04d}"

        if key in seen_activity_keys:
            raise WorkbookMigrationError(f"Duplicate Activity_Key in source sheet: {key}")
        seen_activity_keys.add(key)

        activity_record: dict[str, str] = {header: "" for header in ACTIVITY_HEADERS}
        activity_record["Activity_Key"] = key

        for normalized, idx in normalized_headers.items():
            if idx >= len(row_values):
                continue
            value = row_values[idx]
            if normalized == "activity_key":
                continue
            if normalized == "final_action":
                activity_record["Final Action"] = value
                continue

            for target_header in ACTIVITY_HEADERS:
                if _normalize_header(target_header) == normalized:
                    activity_record[target_header] = value
                    break

        if not activity_record["Final Action"]:
            activity_record["Final Action"] = "Save"

        activity_records.append(activity_record)

        output_type = normalize_output_type(activity_record.get("activity_output_type", ""))
        has_inline_training = any(
            row_values[idx] for idx in training_source_map.values() if idx < len(row_values)
        )
        should_add_training = output_type == "training" or has_inline_training
        if should_add_training:
            training_record = {
                "Activity_Key": key,
                "Training Category": "",
                "Organized By": "",
                "Subject of Training": "",
                "Village": "",
                "Amount": "",
                "Total Trainees": "",
                "Total Duration": "",
            }
            for src, dst in INLINE_TRAINING_SOURCE_HEADERS.items():
                idx = training_source_map.get(src)
                if idx is not None and idx < len(row_values):
                    training_record[dst] = row_values[idx]

            if key in seen_training_keys:
                raise WorkbookMigrationError(f"Duplicate Training output data for Activity_Key {key}")
            seen_training_keys.add(key)
            training_records.append(training_record)

    if ACTIVITY_SHEET_NAME in wb.sheetnames:
        del wb[ACTIVITY_SHEET_NAME]
    activities_ws = wb.create_sheet(ACTIVITY_SHEET_NAME, 0)
    created_sheets.append(ACTIVITY_SHEET_NAME)

    activities_ws.append(ACTIVITY_HEADERS)
    for record in activity_records:
        activities_ws.append([record[h] for h in ACTIVITY_HEADERS])

    if LEGACY_ACTIVITY_SHEET_NAME in wb.sheetnames and LEGACY_ACTIVITY_SHEET_NAME != ACTIVITY_SHEET_NAME:
        del wb[LEGACY_ACTIVITY_SHEET_NAME]

    if "Training" in wb.sheetnames:
        del wb["Training"]
    training_ws = wb.create_sheet("Training")
    created_sheets.append("Training")
    training_ws.append(TRAINING_HEADERS)
    for record in training_records:
        training_ws.append([record[h] for h in TRAINING_HEADERS])

    for sheet_name in [
        ACTIVITY_OUTPUT_DEFINITIONS["community_service"].sheet_name,
        ACTIVITY_OUTPUT_DEFINITIONS["beneficiaries"].sheet_name,
        ACTIVITY_OUTPUT_DEFINITIONS["asset"].sheet_name,
    ]:
        if sheet_name not in wb.sheetnames:
            ws = wb.create_sheet(sheet_name)
            ws.append(["Activity_Key"])
            created_sheets.append(sheet_name)

    _validate_post_migration(activity_records, training_records)

    wb.save(path)

    return MigrationReport(
        backup_path=backup,
        activity_rows=len(activity_records),
        training_rows=len(training_records),
        migrated_columns=list(INLINE_TRAINING_SOURCE_HEADERS.keys()),
        created_sheets=created_sheets,
    )


def _validate_post_migration(activity_records: list[dict[str, str]], training_records: list[dict[str, str]]) -> None:
    activity_keys = {normalize_activity_key(row["Activity_Key"]) for row in activity_records}
    if "" in activity_keys:
        raise WorkbookMigrationError("Missing Activity_Key after migration")

    training_index: dict[str, int] = {}
    for row in training_records:
        key = normalize_activity_key(row["Activity_Key"])
        training_index[key] = training_index.get(key, 0) + 1

    training_activity_keys = {
        normalize_activity_key(row["Activity_Key"])
        for row in activity_records
        if normalize_output_type(row.get("activity_output_type", "")) == "training"
    }

    for key in training_activity_keys:
        if not key:
            raise WorkbookMigrationError("Training activity without Activity_Key")
        if training_index.get(key, 0) != 1:
            raise WorkbookMigrationError(
                f"Training activity must have exactly one Training row for Activity_Key {key}"
            )

    orphan_training = [key for key in training_index if key not in activity_keys]
    if orphan_training:
        raise WorkbookMigrationError(
            f"Orphan Training rows found for Activity_Key: {', '.join(orphan_training)}"
        )


def main() -> None:
    report = migrate_input_workbook("data/input.xlsx")
    print(f"Backup created: {report.backup_path}")
    print(f"Activities rows: {report.activity_rows}")
    print(f"Training rows: {report.training_rows}")
    print(f"Sheets created: {', '.join(report.created_sheets)}")


if __name__ == "__main__":
    main()
