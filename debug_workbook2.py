from openpyxl import load_workbook

wb = load_workbook(r'D:\Python Automation\egram_bot\data\input.xlsx')

print("=== Activities Sheet - ALL COLUMNS ===")
acts = wb['Activities']
print(f"Max column: {acts.max_column}, Max row: {acts.max_row}")
print("\nAll headers:")
headers = [cell.value for cell in acts[1]]
print(headers)

print("\nRow 2 (ACT-0122):")
row2 = [cell.value for cell in acts[2]]
print(row2)

print("\n=== Asset Sheet - ALL DATA ===")
asset = wb['Asset']
print(f"Max column: {asset.max_column}, Max row: {asset.max_row}")
headers_asset = [cell.value for cell in asset[1]]
print("Headers:", headers_asset)

print("\nAll data rows:")
for row_num in range(2, asset.max_row + 1):
    row_data = [cell.value for cell in asset[row_num]]
    print(f"Row {row_num}: {row_data}")

print("\n=== Checking for ACT-0122 in Asset ===")
for row_num in range(2, asset.max_row + 1):
    cell_val = asset[row_num][0].value
    print(f"Row {row_num} Activity_Key: {cell_val} (type: {type(cell_val).__name__})")
